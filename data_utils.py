# =============================================================
# data_utils.py — Calculos sismicos (Norma E.030) y utilidades
# de procesamiento de datos.  Sin dependencias de Streamlit ni
# de Matplotlib.
# =============================================================

import re
import io
import numpy as np
import pandas as pd

# -------------------------------------------------------------
# ALIAS Y CONSTANTES COMPARTIDAS
# -------------------------------------------------------------
VALORES_PREDETERMINADOS = {
    "alturas_input": "3.0, 3.0, 3.0, 3.0, 3.0",
    "periodos_input": "0.556, 0.193, 0.126",
    "masas_input":    "31.8, 31.8, 31.8, 31.8, 27.5",
    "f1_input": "0.03112, 0.05959, 0.08302, 0.09940, 0.10834",
    "f2_input": "-0.08102, -0.10193, -0.05484, 0.03358, 0.10609",
    "f3_input": "0.10415, 0.03021, -0.09525, -0.05769, 0.09176",
}

ALIAS_PARAMETROS = {
    "alturas_entrepiso": "alturas_input", "alturas": "alturas_input",
    "alturas_input": "alturas_input", "altura_entrepiso": "alturas_input",
    "altura_de_entrepiso": "alturas_input", "alturas_de_entrepiso": "alturas_input",
    "alturas_entrepiso_m": "alturas_input", "alturas_de_entrepiso_m": "alturas_input",
    "altura_piso": "alturas_input",
    "periodos_modales": "periodos_input", "periodos": "periodos_input",
    "periodos_input": "periodos_input", "periodo_modal": "periodos_input",
    "periodos_modales_s": "periodos_input", "periodos_s": "periodos_input",
    "vector_masas_piso": "masas_input", "masas": "masas_input",
    "masas_input": "masas_input", "masas_de_pisos": "masas_input",
    "masa_pisos": "masas_input", "masas_pisos": "masas_input",
    "vector_masas": "masas_input", "masas_de_pisos_tonf_s2_m": "masas_input",
    "forma_modal_1": "f1_input", "f1": "f1_input", "f1_input": "f1_input",
    "forma_modal_1_vector": "f1_input", "vector_modal_1": "f1_input",
    "forma_modal_2": "f2_input", "f2": "f2_input", "f2_input": "f2_input",
    "forma_modal_2_vector": "f2_input", "vector_modal_2": "f2_input",
    "forma_modal_3": "f3_input", "f3": "f3_input", "f3_input": "f3_input",
    "forma_modal_3_vector": "f3_input", "vector_modal_3": "f3_input",
    "parametro": "__parametro__", "parametros": "__parametro__",
    "parameter": "__parametro__", "nombre": "__parametro__",
    "valores": "__valores__", "valor": "__valores__",
    "value": "__valores__", "values": "__valores__",
}

_CLAVES_DESTINO = {"alturas_input", "periodos_input", "masas_input",
                   "f1_input", "f2_input", "f3_input"}


# -------------------------------------------------------------
# UTILIDADES AUXILIARES
# -------------------------------------------------------------

def es_numero_simple(valor):
    """Devuelve True si valor puede convertirse a float."""
    try:
        float(str(valor).replace(",", "."))
        return True
    except (TypeError, ValueError):
        return False


def normalizar_nombre(valor):
    """Normaliza un nombre de parametro: minusculas, sin tildes, sin unidades."""
    texto = str(valor).strip().lower()
    traduccion = str.maketrans({"a": "a", "e": "e", "i": "i",
                                "o": "o", "u": "u", "u": "u", "n": "n",
                                "\xe1": "a", "\xe9": "e", "\xed": "i",
                                "\xf3": "o", "\xfa": "u", "\xfc": "u",
                                "\xf1": "n"})
    texto = texto.translate(traduccion)
    texto = re.sub(r"\([^)]*\)", "", texto)
    texto = re.sub(r"[^a-z0-9]+", "_", texto).strip("_")
    return texto


# -------------------------------------------------------------
# CARGA Y PARSEO DE ARCHIVOS DE PARAMETROS
# -------------------------------------------------------------

def leer_archivo_parametros(bytes_archivo, nombre_archivo):
    """Lee un Excel, CSV o TXT y devuelve un dict {clave_destino: texto_csv}."""

    extension = nombre_archivo.rsplit(".", 1)[-1].lower()

    if extension in ("xlsx", "xls"):
        try:
            libro = pd.ExcelFile(io.BytesIO(bytes_archivo))
            tablas = [pd.read_excel(libro, sheet_name=h) for h in libro.sheet_names]
            tabla_carga = next((t for t in tablas if not t.dropna(how="all").empty), tablas[0])
            if tabla_carga.empty or not any(
                ALIAS_PARAMETROS.get(normalizar_nombre(c)) == "__parametro__"
                for c in tabla_carga.columns
            ):
                tablas_nh = [pd.read_excel(libro, sheet_name=h, header=None)
                             for h in libro.sheet_names]
                candidatas = [t for t in tablas_nh if not t.dropna(how="all").empty]
                if candidatas:
                    tabla_carga = candidatas[0]
        except Exception:
            tabla_carga = pd.read_excel(io.BytesIO(bytes_archivo), header=None)
    elif extension == "csv":
        try:
            tabla_carga = pd.read_csv(
                io.BytesIO(bytes_archivo), sep=None, engine="python", encoding="utf-8-sig"
            )
        except Exception:
            tabla_carga = pd.read_csv(
                io.BytesIO(bytes_archivo), header=None, encoding="utf-8-sig"
            )
    else:
        contenido = bytes_archivo.decode("utf-8-sig")
        try:
            tabla_carga = pd.read_csv(
                io.StringIO(contenido), sep=None, engine="python", comment="#"
            )
        except Exception:
            filas = [linea.split("=", 1)
                     for linea in contenido.splitlines() if "=" in linea]
            tabla_carga = pd.DataFrame(filas, columns=["Parametro", "Valores"])

    tabla_carga = tabla_carga.dropna(how="all").dropna(axis=1, how="all")
    pares = []

    if not tabla_carga.empty:
        columnas_norm = [normalizar_nombre(c) for c in tabla_carga.columns]
        idx_nombre = next(
            (i for i, c in enumerate(columnas_norm)
             if ALIAS_PARAMETROS.get(c) == "__parametro__"), None
        )
        idx_valor = next(
            (i for i, c in enumerate(columnas_norm)
             if ALIAS_PARAMETROS.get(c) == "__valores__"), None
        )
        if idx_nombre is not None and idx_valor is not None:
            for _, fila in tabla_carga.iterrows():
                n, v = fila.iloc[idx_nombre], fila.iloc[idx_valor]
                if pd.notna(n) and pd.notna(v):
                    pares.append((n, v))
        else:
            for _, fila in tabla_carga.iterrows():
                celdas = list(fila)
                if not celdas or pd.isna(celdas[0]):
                    continue
                nom_norm = normalizar_nombre(celdas[0])
                if ALIAS_PARAMETROS.get(nom_norm) in _CLAVES_DESTINO:
                    vals = [v for v in celdas[1:] if pd.notna(v)]
                    if vals:
                        pares.append((celdas[0], vals[0] if len(vals) == 1 else vals))
            if not pares:
                for i, col in enumerate(tabla_carga.columns):
                    destino = ALIAS_PARAMETROS.get(normalizar_nombre(col))
                    if destino in _CLAVES_DESTINO:
                        vals = [v for v in tabla_carga.iloc[:, i].tolist() if pd.notna(v)]
                        if vals:
                            pares.append((col, vals if len(vals) > 1 else vals[0]))

    resultado = {}
    for nombre, valor in pares:
        destino = ALIAS_PARAMETROS.get(normalizar_nombre(nombre))
        if destino not in _CLAVES_DESTINO:
            continue
        if valor is None or (isinstance(valor, (float, np.floating)) and pd.isna(valor)):
            continue
        if isinstance(valor, (list, tuple, np.ndarray, pd.Series)):
            texto = ", ".join(str(x).strip() for x in list(valor) if pd.notna(x))
        else:
            texto = str(valor).strip()
        texto = texto.replace(";", ",")
        if "," not in texto and " " in texto.strip():
            partes = texto.split()
            if len(partes) > 1 and all(es_numero_simple(x) for x in partes):
                texto = ", ".join(partes)
        resultado[destino] = texto

    return resultado


# -------------------------------------------------------------
# PARSEO DE ENTRADAS DE TEXTO
# -------------------------------------------------------------

def parsear_entradas(alturas_txt, periodos_txt, masas_txt, f1_txt, f2_txt, f3_txt):
    """Convierte textos CSV de la UI en arrays numpy.

    Devuelve:
        alturas_entrepiso, periodos_modales, vector_masas_piso,
        forma_modal_1, forma_modal_2, forma_modal_3,
        matriz_masas, matriz_formas_modales, num_pisos, lista_pisos, Ht
    """
    alturas_entrepiso    = np.array([float(x.strip()) for x in alturas_txt.split(",")])
    periodos_modales     = np.array([float(x.strip()) for x in periodos_txt.split(",")])
    vector_masas_piso    = np.array([float(x.strip()) for x in masas_txt.split(",")])
    forma_modal_1        = np.array([float(x.strip()) for x in f1_txt.split(",")])
    forma_modal_2        = np.array([float(x.strip()) for x in f2_txt.split(",")])
    forma_modal_3        = np.array([float(x.strip()) for x in f3_txt.split(",")])

    matriz_masas          = np.diag(vector_masas_piso)
    matriz_formas_modales = [forma_modal_1, forma_modal_2, forma_modal_3]
    num_pisos             = len(alturas_entrepiso)
    lista_pisos           = [f"Piso {i+1}" for i in range(num_pisos)]
    Ht                    = float(np.sum(alturas_entrepiso))

    return (alturas_entrepiso, periodos_modales, vector_masas_piso,
            forma_modal_1, forma_modal_2, forma_modal_3,
            matriz_masas, matriz_formas_modales, num_pisos, lista_pisos, Ht)


def validar_entradas(periodos_modales, vector_masas_piso, num_pisos,
                     alturas_entrepiso, matriz_formas_modales):
    """Devuelve una lista de mensajes de error (vacia si todo esta bien)."""
    errores = []
    if len(periodos_modales) != 3:
        errores.append("Ingresa exactamente 3 periodos modales.")
    if len(vector_masas_piso) != num_pisos:
        errores.append(
            f"Hay {num_pisos} alturas de entrepiso pero {len(vector_masas_piso)} masas."
        )
    if any(len(f) != num_pisos for f in matriz_formas_modales):
        errores.append("Cada forma modal debe tener un valor por piso.")
    if (np.any(alturas_entrepiso <= 0) or np.any(vector_masas_piso <= 0)
            or np.any(periodos_modales <= 0)):
        errores.append("Alturas, masas y periodos deben ser mayores que cero.")
    return errores


# -------------------------------------------------------------
# FUNCIONES DE CALCULO SISMICO (E.030)
# -------------------------------------------------------------

def calcular_parametros_espectrales(Z, U, S, R, Tp, TL, Ht, Ct, periodos_modales):
    """Calcula periodo fundamental, C, Sa, omega y Sd para cada modo."""
    g = 9.81
    T = Ht / Ct
    C = np.array([
        2.5 if p < Tp
        else (2.5 * (Tp / p) if p <= TL else 2.5 * (Tp * TL / p**2))
        for p in periodos_modales
    ])
    Sa = (Z * U * C * S / R) * g
    w  = 2 * np.pi / periodos_modales
    Sd = Sa / (w**2)
    return T, C, Sa, w, Sd


def calcular_factores_participacion(matriz_masas, vector_masas_piso, matriz_formas_modales):
    """Calcula los factores de participacion modal gamma para cada modo."""
    gammas = []
    for phi in matriz_formas_modales:
        num = np.dot(phi, vector_masas_piso)
        den = np.dot(phi, np.dot(matriz_masas, phi))
        gammas.append(num / den)
    return np.array(gammas)


def calcular_respuestas_modales(vector_masas_piso, matriz_formas_modales, Sa, gammas):
    """Calcula fuerzas y cortantes modales por piso para cada modo."""
    fuerzas, cortantes = [], []
    for k, phi in enumerate(matriz_formas_modales):
        d = Sa[k] * gammas[k] * phi
        f = vector_masas_piso * d
        v = np.cumsum(f[::-1])[::-1]
        fuerzas.append(f)
        cortantes.append(v)
    return np.array(fuerzas), np.array(cortantes)


def calcular_combinacion_modal(cortantes_modales, R, periodos_modales, amortiguamiento=0.05):
    """Combina respuestas modales por ABS, SRSS y CQC.

    El CQC utiliza coeficientes de correlacion entre modos para una razon de
    amortiguamiento (por defecto 5%). Las respuestas modales deben conservar
    su signo para aplicar la expresion cuadratica completa.

    Devuelve:
        V_abs, V_srss, V_cqc, V_real
    """
    V_abs  = np.sum(np.abs(cortantes_modales), axis=0)
    V_srss = np.sqrt(np.sum(cortantes_modales**2, axis=0))

    freqs = 2 * np.pi / np.asarray(periodos_modales, dtype=float)
    n = len(freqs)
    rho = np.eye(n)
    for i in range(n):
        for j in range(i + 1, n):
            b   = min(freqs[i], freqs[j]) / max(freqs[i], freqs[j])
            num = 8 * amortiguamiento**2 * (1 + b) * b**1.5
            den = (1 - b**2)**2 + 4 * amortiguamiento**2 * b * (1 + b)**2
            r   = num / den if den > 0 else 0.0
            rho[i, j] = rho[j, i] = r

    V_cqc = np.zeros(cortantes_modales.shape[1])
    for piso in range(cortantes_modales.shape[1]):
        r = cortantes_modales[:, piso]
        V_cqc[piso] = np.sqrt(max(float(r @ rho @ r), 0.0))

    # Factor de escalamiento segun criterio de diseno del proyecto.
    # Verificar su correspondencia con el criterio exigido en clase.
    V_real = V_cqc * 0.75 * R
    return V_abs, V_srss, V_cqc, V_real


# -------------------------------------------------------------
# GENERACION DE RESULTADOS TABULARES
# -------------------------------------------------------------

def generar_dataframe_resultados(lista_pisos, fuerzas_modales, cortantes_modales,
                                 V_abs, V_srss, V_cqc, V_real):
    """Ensambla el DataFrame principal de resultados."""
    return pd.DataFrame({
        "Piso":                lista_pisos,
        "Fuerza_Modo_1":       fuerzas_modales[0],
        "Fuerza_Modo_2":       fuerzas_modales[1],
        "Fuerza_Modo_3":       fuerzas_modales[2],
        "Cortante_Modo_1":     cortantes_modales[0],
        "Cortante_Modo_2":     cortantes_modales[1],
        "Cortante_Modo_3":     cortantes_modales[2],
        "Cortante_Suma_Abs":   V_abs,
        "Cortante_SRSS":       V_srss,
        "Cortante_CQC":        V_cqc,
        "Cortante_Real_Diseno": V_real,
    })


def estilar_dataframe(df, es_oscuro):
    """Aplica estilos CSS al DataFrame segun el tema activo."""
    bg_hdr    = "#0284c7"
    text_hdr  = "#ffffff"
    bg_cell   = "#1e293b" if es_oscuro else "#ffffff"
    text_cell = "#f8fafc" if es_oscuro else "#1e293b"
    border    = "rgba(255,255,255,0.15)" if es_oscuro else "#cbd5e1"
    bg_alt    = "#0f172a" if es_oscuro else "#f8fafc"

    fmt = {col: "{:.3f}" for col in df.columns if col != "Piso"}
    styles = [
        {"selector": "th", "props": [
            ("background-color", bg_hdr), ("color", text_hdr),
            ("font-weight", "bold"), ("font-family", "Times New Roman"),
            ("border", f"1px solid {border}"), ("text-align", "center"),
            ("padding", "10px 12px"), ("font-size", "13px"),
        ]},
        {"selector": "td", "props": [
            ("background-color", bg_cell), ("color", text_cell),
            ("font-family", "Times New Roman"), ("border", f"1px solid {border}"),
            ("text-align", "center"), ("padding", "8px 12px"), ("font-size", "13px"),
        ]},
        {"selector": "table", "props": [
            ("border-collapse", "collapse"), ("width", "100%"),
            ("border-radius", "8px"), ("overflow", "hidden"),
            ("box-shadow", "0 2px 8px rgba(0,0,0,0.06)"),
        ]},
        {"selector": "tbody tr:nth-child(even) td", "props": [
            ("background-color", bg_alt),
        ]},
    ]
    return df.style.format(fmt).set_table_styles(styles)


# -------------------------------------------------------------
# EXPORTACION A EXCEL
# -------------------------------------------------------------

def generar_excel_bytes(df):
    """Genera un reporte .xlsx formateado y devuelve el buffer BytesIO."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte_E030"

    font_titulo = Font(name="Times New Roman", size=13, bold=True, color="0284C7")
    font_header = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
    font_body   = Font(name="Times New Roman", size=10)
    fill_header = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    fill_alt    = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    align_c     = Alignment(horizontal="center", vertical="center")
    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.append(["REPORTE DE ANALISIS SISMICO MODAL ESPECTRAL - NORMA E.030"])
    ws.cell(row=1, column=1).font = font_titulo
    ws.append([])
    ws.append(list(df.columns))
    for col_num in range(1, len(df.columns) + 1):
        cell = ws.cell(row=3, column=col_num)
        cell.font      = font_header
        cell.fill      = fill_header
        cell.alignment = align_c
        cell.border    = border_thin

    for row_idx, row_val in enumerate(df.values, start=4):
        ws.append(list(row_val))
        for col_num in range(1, len(row_val) + 1):
            cell           = ws.cell(row=row_idx, column=col_num)
            cell.font      = font_body
            cell.alignment = align_c
            cell.border    = border_thin
            if row_idx % 2 == 0:
                cell.fill = fill_alt
            if isinstance(cell.value, (float, int)) and col_num > 1:
                cell.number_format = "0.000"

    for col in ws.columns:
        max_len = max(
            (len(str(c.value)) for c in col if c.row != 1 and c.value is not None),
            default=0,
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 5, 14)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
